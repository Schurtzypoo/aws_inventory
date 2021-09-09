import json
import boto3
import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from botocore.config import Config


class aws_inventory:
    def __init__(self):
        return
    
    def get_instances(self, ec2_config):
        instanceData = ec2_config.describe_instances()
        #print(instanceData)
        return instanceData

    def get_vpcs(self, ec2_config):
        vpcData = ec2_config.describe_vpcs()
        # print(vpcData)
        return vpcData

    def get_subnets(self, ec2_config):
        subnetData = ec2_config.describe_subnets()
        # print(subnetData)
        return subnetData

    def build_dict(self, ec2_config):
        dataDict = {
            'Instance Data': self.get_instances(ec2_config),
            'VPC Data': self.get_vpcs(ec2_config),
            'Subnet Data': self.get_subnets(ec2_config)
            }
        return dataDict
    
    def get_instance_tags(self, instanceData):
        required_tags = ["FISMA ID", "Name", "Operating_System", "Account", "Purpose"]
        tagList = []
        for item in required_tags:
            tagList.append("")
        for tag in instanceData["Tags"]:
            for index, item in enumerate(required_tags):
                if tag["Key"] == item:
                    tagList[index] = tag["Value"]
        return tagList

    def get_instance_subnet(self, instanceData, subnetDict):
        required_info = ["Subnet ID", "Subnet", "Subnet CIDR", "Availability Zone"]
        subnetInfo = []
        for item in required_info:
            subnetInfo.append("")
        subnetInfo[0] = instanceData["SubnetId"]
        for subnet in subnetDict["Subnets"]:
            if subnet["SubnetId"] == subnetInfo[0]:
                subnetInfo[2] = subnet["CidrBlock"]
                subnetInfo[3] = subnet["AvailabilityZone"]
                for tag in subnet["Tags"]:
                    if tag["Key"] == "Name":
                        subnetInfo[1] = tag["Value"]

        return subnetInfo
    
    def get_vpc_info(self, vpcData, vpcId):
        for vpc in vpcData["Vpcs"]:
            if vpcId == vpc["VpcId"]:
                for tag in vpc["Tags"]:
                    if tag["Key"] == "Name":
                        vpcName = tag["Value"]
        return vpcName
    
    def get_nic_info(self, instanceData):
        nicCount = len(instanceData)
        return nicCount
    
    def get_volume_info(self, instanceData):
        volumeCount = len(instanceData["BlockDeviceMappings"])
        return volumeCount
    
    def get_ip_info(self, interfaceData):
        required_ips = ["Private Management", "Public Management"]
        ipInfo = []
        for item in required_ips:
            ipInfo.append("")
        for interface in interfaceData:
            description = interface["Description"]
            description = description.lower()
            if ("management" or "mgmt") in description:
                ipInfo[0] = interface["PrivateIpAddresses"][0]["PrivateIpAddress"]
                try:
                    ipInfo[1] = interface["PrivateIpAddresses"][0]["Association"]["PublicIp"]
                except KeyError:
                    pass
                break
            else:
                ipInfo[0] = interface["PrivateIpAddress"]
                try:
                    ipInfo[1] = interface["Association"]["PublicIp"]
                except KeyError:
                    pass
        return ipInfo

    def populate_ec2_data(self, dataDict, ec2_dataList, region):
        
        data = dataDict["Instance Data"]
        for object in data["Reservations"]:
            instanceData = object["Instances"][0]
            if instanceData["State"]["Name"] != "terminated":
                state = instanceData["State"]["Name"]
                vpcId = instanceData["VpcId"]
                vpcName = self.get_vpc_info(dataDict["VPC Data"], vpcId)
                instanceId = instanceData["InstanceId"]
                amiId = instanceData["ImageId"]
                ipInfo = self.get_ip_info(instanceData["NetworkInterfaces"])
                instanceRegion = region
                tagList = self.get_instance_tags(instanceData)
                subnetInfo = self.get_instance_subnet(instanceData, dataDict["Subnet Data"])
                nicCount = self.get_nic_info(instanceData["NetworkInterfaces"])
                volumeCount = self.get_volume_info(instanceData)
                ec2_dataList.append([instanceRegion, tagList[1], state, tagList[3], tagList[4], instanceId, tagList[0], ipInfo[0], ipInfo[1], nicCount, volumeCount, amiId, tagList[2], vpcName, vpcId, subnetInfo[1], subnetInfo[2], subnetInfo[3], subnetInfo[0]])

    def write_ec2_data(self, ec2_dataList, ws):
        for row in ec2_dataList:
            ws.append(row) 
        table = Table(displayName="EC2Data", ref=f"A1:{chr(ord('@')+len(ec2_dataList[0]))}{len(ec2_dataList)}")
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        table.tableStyleInfo = style
        ws.add_table(table)

    def make_worksheet(self, workbook, ws_title, index):
        ws = workbook.create_sheet(title=ws_title, index=index)
        ws = workbook.active
        return ws

    def output_json(self, ec2_config, region):
        dataDict = self.build_dict(ec2_config)
        for k,v in dataDict.items():
            with open(f'{region}.{k}.json','w+') as f:
                json.dump(v, f, indent = 4, default=str)
                f.close()

    def compile(self):
        workbook = Workbook()
        ec2_dataList = [["Region", "Instance Name", "State", "Account Name", "Purpose", "Instance ID", "FISMA ID", "Private IP", "Public IP", "NIC Count", "Volume Count", "AMI ID", "Operating System", "VPC", "VPC ID", "Subnet", "Subnet CIDR", "Availability Zone", "Subnet ID"]]
        ec2_ws = self.make_worksheet(workbook, "EC2 Data", 0)
        for region in ['us-gov-east-1', "us-gov-west-1"]:
            sub_config = Config(
                region_name = region
            )
            ec2_config = boto3.client('ec2', config=sub_config)
            dataDict = self.build_dict(ec2_config)
            self.populate_ec2_data(dataDict, ec2_dataList, region)
            #self.output_json(ec2_config, region)
        self.write_ec2_data(ec2_dataList, ec2_ws)
        workbook.save(filename="aws_inventory.xlsx")

def main():
    inventory = aws_inventory()
    inventory.compile()

if __name__ == "__main__":
    main()
